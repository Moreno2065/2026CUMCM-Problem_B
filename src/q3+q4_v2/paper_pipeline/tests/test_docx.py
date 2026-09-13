from pathlib import Path

from docx import Document

from paper_pipeline.update_paper_docx import update


def test_final_docx_updates_symbols_and_orders_figures(tmp_path: Path):
    project = Path(__file__).resolve().parents[4]
    source = project / "paper" / "B题论文(1).docx"
    evidence = project / "src" / "q3+q4_v2" / "paper_evidence"
    output = tmp_path / "paper.docx"

    report = update(source, output, evidence)
    assert report["q3_q4_symbol_rows_updated"] == 9
    assert report["figures_embedded"] == 9

    doc = Document(output)
    symbols = doc.tables[0]
    expected = {
        23: ("X_c", "频道 c 的完整物理可行状态集合", ""),
        33: ("H_c⁺", "频道 c 已获得正观测点的安全凸包", ""),
        34: ("N_c", "频道 c 的无信号见证点集合", ""),
        35: ("S₂₁", "问题四的 21 站连续覆盖证书集合", ""),
        36: ("rᵢₙ, rₒᵤₜ", "21 站证书网的内、外环半径", "米"),
        37: ("x₊, x₋", "crossbar 截面两侧的条件探点", "米"),
        38: ("h_cb", "crossbar 条件探点的横向偏移量", "米"),
        39: ("K_rot", "初始站网旋转候选数", "个"),
        40: ("N_max", "每频道无线测量次数上限", "次"),
    }
    for index, values in expected.items():
        assert tuple(cell.text for cell in symbols.rows[index].cells) == values

    captions = [
        paragraph.text.strip()
        for paragraph in doc.paragraphs
        if paragraph.style.name == "图表标题"
        and paragraph.text.strip().startswith("图 ")
    ]
    q34_captions = [caption for caption in captions if caption.split()[1].isdigit()
                    and 7 <= int(caption.split()[1]) <= 15]
    assert [int(caption.split()[1]) for caption in q34_captions] == list(range(7, 16))
