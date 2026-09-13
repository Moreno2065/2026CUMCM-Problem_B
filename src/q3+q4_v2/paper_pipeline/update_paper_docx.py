"""Replace Q3/Q4 manuscript sections with the verified V1/V2 paper narrative.

The original DOCX is never overwritten.  Q1/Q2 model XML is hashed before and
after editing so an accidental cross-section mutation fails the build.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree


def _find_paragraph(doc: Document, exact: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == exact:
            return paragraph
    raise ValueError(f"paragraph not found: {exact}")


def _find_prefix(doc: Document, prefix: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"paragraph prefix not found: {prefix}")


def _paragraph_index(doc: Document, target) -> int:
    return next(i for i, paragraph in enumerate(doc.paragraphs)
                if paragraph._p is target._p)


def _remove_between(start, end, include_start: bool = False) -> None:
    parent = start._p.getparent()
    nodes = list(parent)
    i0 = nodes.index(start._p) + (0 if include_start else 1)
    i1 = nodes.index(end._p)
    for node in nodes[i0:i1]:
        parent.remove(node)


def _move_before(node, anchor) -> None:
    parent = anchor._p.getparent()
    parent.insert(parent.index(anchor._p), node)


def _set_run_font(run, *, size: float = 10.5, bold: bool | None = None,
                  east_asia: str = "宋体", latin: str = "Times New Roman") -> None:
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def _format_body(paragraph) -> None:
    paragraph.paragraph_format.first_line_indent = Pt(21)
    paragraph.paragraph_format.line_spacing = 1.35
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in paragraph.runs:
        _set_run_font(run)


def _add_paragraph_before(doc: Document, anchor, text: str,
                          style: str = "Normal", first_indent: bool = True,
                          keep_next: bool | None = None):
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    if style.startswith("Heading"):
        _set_run_font(run, size=12 if style == "Heading 2" else 10.5,
                      bold=True, east_asia="黑体")
        paragraph.paragraph_format.keep_with_next = True
    elif style == "图表标题":
        _set_run_font(run, size=9, east_asia="宋体")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = bool(keep_next)
    else:
        _format_body(paragraph)
        if not first_indent:
            paragraph.paragraph_format.first_line_indent = Pt(0)
    _move_before(paragraph._p, anchor)
    return paragraph


def _add_equation_before(doc: Document, anchor, text: str):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    _set_run_font(run, size=10.5, east_asia="Cambria Math", latin="Cambria Math")
    _move_before(paragraph._p, anchor)
    return paragraph


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _add_table_before(doc: Document, anchor, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    _set_repeat_table_header(table.rows[0])
    for j, value in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for values in rows:
        cells = table.add_row().cells
        for j, value in enumerate(values):
            cells[j].text = str(value)
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    _set_run_font(run, size=8.5, bold=(row_index == 0),
                                  east_asia="宋体")
                if row_index < len(table.rows) - 1:
                    paragraph.paragraph_format.keep_with_next = True
    _move_before(table._tbl, anchor)
    return table


def _add_picture_before(doc: Document, anchor, path: Path, width: float = 6.15):
    doc.add_picture(str(path), width=Inches(width))
    paragraph = doc.paragraphs[-1]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    _move_before(paragraph._p, anchor)


def _range_hash(doc: Document, start_text: str, end_text: str) -> str:
    start = _find_paragraph(doc, start_text)._p
    end = _find_paragraph(doc, end_text)._p
    parent = start.getparent()
    nodes = list(parent)
    payload = b"".join(etree.tostring(node, encoding="utf-8")
                       for node in nodes[nodes.index(start):nodes.index(end)])
    return hashlib.sha256(payload).hexdigest()


def _read_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _paired_rows(evidence: Path, question: str) -> tuple[list[list[str]], dict]:
    rows = _read_csv(evidence / "tables/paired_effects.csv")
    selected = [row for row in rows if row["question"] == question]
    all_row = next(row for row in selected if row["source_count"] == "all")
    table = []
    for count in (10, 13, 16):
        row = next(row for row in selected if row["source_count"] == str(count))
        table.append([
            str(count), row["n_pairs"], f'{float(row["mean_v1_s"]):.2f}',
            f'{float(row["mean_v2_s"]):.2f}',
            f'{float(row["relative_reduction_pct"]):.1f}%',
            f'{row["wins_v2"]}/{row["n_pairs"]}',
        ])
    return table, all_row


def _replace_short_sections(doc: Document) -> None:
    q3_abs = _find_prefix(doc, "针对问题三，")
    q3_abs.text = (
        "针对问题三，先以集合知识状态、有限发现与清除证书、基数停止规则构造 V1 "
        "可执行基线，再在同一完整 verifier 门槛下建立 V2 联合搜索—定位—清除策略。"
        "在 60 个共同案例上，V2 相对 V1 的平均每源虚拟时间降低 48.0%（95% bootstrap "
        "CI：46.6%～49.3%），60/60 个案例更快；消融表明主要节时来自联合路线。"
    )
    _format_body(q3_abs)
    q4_abs = _find_prefix(doc, "针对问题四，")
    q4_abs.text = (
        "针对问题四，V2 用中心点、8 点内环和 12 点外环组成的 21 站连续覆盖证书处理"
        "定向源，并以 crossbar 双侧条件探点消解无信号的方向歧义。在 60 个 mixed "
        "共同案例上，V2 平均每源虚拟时间降低 59.4%（95% bootstrap CI：57.9%～60.9%），"
        "60/60 个案例更快。去除 16 源基数停止后，16 源案例仅 3/20 通过完整 verifier，"
        "说明该规则是终止证据的一部分。两问共 240 次主运行均完成全部清除、通过 verifier "
        "且计时账本闭合。"
    )
    _format_body(q4_abs)
    keywords = _find_prefix(doc, "关键词：")
    keywords.text = (
        "关键词：有界误差；集合知识状态；确定性证书；联合开放路线；定向源；配对实验"
    )
    _format_body(keywords)

    q3_head = _find_paragraph(doc, "问题三：")
    q4_head = _find_paragraph(doc, "问题四：")
    q3_head.text = "问题三："
    q3_index = _paragraph_index(doc, q3_head)
    q3_a = doc.paragraphs[q3_index + 1]
    q3_b = doc.paragraphs[q3_index + 2]
    q3_a.text = (
        "问题三要求在全向源数量、位置和接收半径均未知的条件下完成全部搜索与清除。"
        "判断一次测量是否有效并不困难，真正的难点是同时处理完成性和效率：每个真实源必须"
        "被发现并清除，无源频道必须有可审计证据，任务还要尽量减少固定扫描、返回锚点和"
        "重复测量的耗时。"
    )
    q3_b.text = (
        "本文把两类要求分层处理。V1 固定为状态—证书—停止的保证性基线；V2 完整保留"
        "这些门禁，只重排合法动作，将扫描、局部定位与清除放进同一条开放路线。"
        "两版在相同案例、相同种子和相同计时规则下配对运行，由此估计决策层优化的净收益。"
    )
    for paragraph in (q3_a, q3_b):
        _format_body(paragraph)

    q4_index = _paragraph_index(doc, q4_head)
    q4_a = doc.paragraphs[q4_index + 1]
    q4_b = doc.paragraphs[q4_index + 2]
    q4_a.text = (
        "问题四加入定向源后，测点返回 no_signal 既可能表示频道不存在，也可能只是位于"
        "发射半平面之外。问题三的全向圆盘证空逻辑不能直接套用，否则会出现漏源和假证空。"
    )
    q4_b.text = (
        "最终模型继续采用 V1 的集合状态和终止门禁，但用经连续几何核验的 21 站证书网"
        "完成发现，并加入 crossbar 双侧条件探点、负区域更新和源任务精确插入。"
        "每次无信号只在满足明确几何条件时收缩可行域，发现骨干和局部服务共同参与路线决策。"
    )
    for paragraph in (q4_a, q4_b):
        _format_body(paragraph)


def _update_q3_q4_symbols(doc: Document) -> None:
    """Normalize Q3/Q4 symbols and replace obsolete V1 implementation rows."""
    symbol_table = next(
        table for table in doc.tables
        if len(table.columns) == 3
        and table.rows
        and table.rows[0].cells[0].text.strip() == "符号"
    )
    replacements = {
        23: ("X_c", "频道 c 的完整物理可行状态集合", ""),
        33: ("H_c⁺", "频道 c 已获得正观测点的安全凸包", ""),
        34: ("N_c", "频道 c 的无信号见证点集合", ""),
        35: ("S₂₁", "问题四的 21 站连续覆盖证书集合", ""),
        36: ("rᵢₙ, rₒᵤₜ", "21 站证书网的内、外环半径", "米"),
        37: ("x₊, x₋", "crossbar 截面两侧的条件探点", "米"),
        38: ("h_cb", "crossbar 条件探点的横向偏移量", "米"),
        39: ("K_rot", "初始站网旋转候选数", "个"),
        40: ("N_max", "每频道无线测量次数上限", "次"),
    }
    for row_index, values in replacements.items():
        row = symbol_table.rows[row_index]
        for cell, value in zip(row.cells, values):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    _set_run_font(run, size=9)


def _build_model_sections(doc: Document, anchor, evidence: Path) -> None:
    figures = evidence / "figures"
    q3_table, q3_all = _paired_rows(evidence, "q3")
    q4_table, q4_all = _paired_rows(evidence, "q4")

    H2 = lambda text: _add_paragraph_before(doc, anchor, text, "Heading 2")
    H3 = lambda text: _add_paragraph_before(doc, anchor, text, "Heading 3")
    P = lambda text: _add_paragraph_before(doc, anchor, text)
    EQ = lambda text: _add_equation_before(doc, anchor, text)
    PIC = lambda name, width=6.15: _add_picture_before(doc, anchor, figures / name, width)
    CAP = lambda text: _add_paragraph_before(doc, anchor, text, "图表标题")
    TCAP = lambda text: _add_paragraph_before(
        doc, anchor, text, "图表标题", keep_next=True)

    H2("问题三模型建立与求解")
    H3("保证性状态与 V1 冻结基线")
    P("对频道 c 定义知识状态 Kc=(Pc, Ec, Oc, sc)，其中 Pc 是由有界测向角域和接收约束得到的位置可行集，Ec 为已证空域，Oc 保存原始观测与清除反馈。状态 sc 取 UNKNOWN、ACTIVE、READY、CLEARED、CERTIFIED_ABSENT 五类。")
    EQ("sc ∈ {UNKNOWN, ACTIVE, READY, CLEARED, CERTIFIED_ABSENT}")
    P("direction 使频道进入 ACTIVE，near 或可行集已落入 20 m 清除圆时进入 READY，清除成功后进入 CLEARED。只有有限发现证书闭合或由源数上界推出时，频道才可进入 CERTIFIED_ABSENT。任务仅在已清除 16 个不同频道，或 20 个频道均为 CLEARED/CERTIFIED_ABSENT 时结束；清除 10 个源不是停止条件。")
    P("V1 采用固定发现点、返回结构和逐频道服务，是满足上述门禁的冻结可执行基线。它用于回答硬误差边界下能否可靠完成，并给 V2 提供不可放松的状态、证书与停止约束。两版使用同一案例和计时规则，V1 同时构成效率比较基线。")
    PIC("fig07_v1_baseline.jpg")
    CAP("图 7  V1 冻结基线的知识状态、有限证书与停止规则")

    H3("V2 联合搜索—定位—清除路线")
    P("设当前位置为 x0，未完成任务集合由扫描、定位和清除任务组成。每个任务 i 有候选位置 xi、动作耗时 τi 和可共享频道集合 Ci。V2 不再依次执行全域扫描、返回和逐源定位，而是在每次重规划时构造从当前位置出发且不要求回到起点的开放路线 π。")
    EQ("𝒯 = 𝒯scan ∪ 𝒯loc ∪ 𝒯clear")
    EQ("J(π) = v⁻¹Σ‖xπ(k+1) − xπ(k)‖₂ + Στπ(k) + Jresidual(π)")
    P("Jresidual 是路线结束后尚未闭合任务的保守代价。算法只比较已经由 V1 保证层判定为合法的动作，因此路线缩短不会改变完成判据。Q3 最终求解器 q3_joint_search_clear_route_v5 使用联合路线、250 m 探点距离、0.7 搜索站共享阈值、40 m 试探清除半径和 1500 m 规划半径。局部定位点投影到开放路线附近并进行一次次序细化；共享测量和试探清除仅在条件满足时触发。")
    P("图 8 给出两版模型的衔接。V1 的集合状态、发现与清除证书、基数停止和 verifier 全部保留，V2 升级的是开放路线、条件探点和局部清除决策。")
    PIC("fig08_framework.jpg")
    CAP("图 8  V1 保证层与 V2 联合决策层的继承关系")

    H3("Q3 算法流程与代表路线")
    for text in (
        "（1）在原点完成当前频道测量，更新 Kc，并建立扫描、定位和清除任务池。",
        "（2）删除已 CLEARED 或具有有效 CERTIFIED_ABSENT 证据的任务，检查 16 源基数停止与 20 频道闭合条件。",
        "（3）计算合法候选点的移动、测量、换频和清除成本，把可在同一位置完成的多频道测量合并。",
        "（4）构造联合开放路线，对局部定位点执行连续投影和一次次序细化，然后执行首个动作并用真实反馈更新状态。",
        "（5）结束后由独立 verifier 复核真实源清除、证空依据、停止条件和虚拟时间账本；任一项失败均不计为可接受运行。",
    ):
        _add_paragraph_before(doc, anchor, text, "List Paragraph", first_indent=False)
    P("代表路线中的真实源坐标只用于离线解释，不参与在线决策。搜索站、定位点和清除点的联合方式见图 14；总体效果仍以全部配对案例为准。")

    H3("Q3 配对实验结果")
    P("V1 与 V2 使用完全相同的案例文件、源位置、半径、频道、种子和计时规则。源数 10、13、16 各取 20 个案例，共 60 对。主指标为每源虚拟时间 tsrc=Tvirtual/Nsource。置信区间以案例为重采样单位，固定种子 20260913，进行 10000 次 percentile bootstrap。只有完成全部真实源清除、通过 verifier 且计时账本闭合的运行进入效率统计。")
    TCAP("表 5  Q3 严格配对结果")
    _add_table_before(doc, anchor,
                      ["源数", "案例对数", "V1/s·源⁻¹", "V2/s·源⁻¹", "平均降幅", "V2 更快"],
                      q3_table)
    P(f'60 对案例中，V1 和 V2 的平均每源时间分别为 {float(q3_all["mean_v1_s"]):.2f} s 和 {float(q3_all["mean_v2_s"]):.2f} s。V2 平均降低 {float(q3_all["relative_reduction_pct"]):.1f}%，95% bootstrap CI 为 {float(q3_all["relative_reduction_ci95_low_pct"]):.1f}%～{float(q3_all["relative_reduction_ci95_high_pct"]):.1f}%，{q3_all["wins_v2"]}/{q3_all["n_pairs"]} 个案例更快。')
    PIC("fig09_paired_performance.jpg")
    CAP("图 9  严格配对案例中的 V1–V2 每源虚拟时间")

    H3("Q3 消融与参数敏感性")
    P("消融实验在 10 源和 16 源案例上各使用 20 个共同案例。去除联合路线后，每源时间分别上升 30.6% 和 29.8%，是 Q3 的主要效率来源；去除站点路线细化后分别上升 9.8% 和 6.1%。共享探点、连续站点投影和试探清除的平均影响接近 0～1%，只能视为条件性补充。全部 240 次 Q3 消融运行均通过 verifier。")
    PIC("fig10_q3_ablation.jpg")
    CAP("图 10  Q3 模块消融的配对每源时间变化")
    P("锁定参数邻域的事后敏感性分析没有用于重新选参。探点距离 150/250/350 m 时平均每源时间为 226.3/223.6/224.3 s，共享阈值 0.5/0.7/0.9 时为 222.3/223.6/225.3 s，试探半径 20/40/60 m 时为 225.6/223.6/223.4 s，规划半径 1000/1500/1800 m 时为 223.2/223.6/224.6 s。锁定值 0.7 并非这批事后案例上的经验最优点。")

    H2("问题四模型建立与求解")
    H3("定向源观测与 21 站连续覆盖证书")
    P("定向源状态写为 (p,r,u)，其中 p 为位置，r 为接收半径，u 为发射朝向单位向量。测点 x 收到信号必须同时满足距离和朝向条件；单点 no_signal 因而不能直接推出频道不存在。")
    EQ("‖x − p‖₂ ≤ r， uᵀ(x − p) ≥ 0")
    P("Q4 最终求解器 q4_21station_exact_route_v4 使用 21 个证书站点：中心 1 点、半径 995 m 的 8 点内环和半径 1864 m 的 12 点外环。程序对连续目标域执行覆盖核验。整张站网的刚性旋转不改变证书；原点首次观测后，算法在 [0,π/2) 内比较 12 个旋转候选，选择预测开放路线最短者。该选择只依赖已取得的观测和任务几何，不读取真实源位置或真实源数。")
    P("一个频道只有在全部必要站点均无信号时才能标为 CERTIFIED_ABSENT。外环的部分测点位于 1800 m 目标圆外，这是连续覆盖构造的一部分，不是越界的真实源位置。")

    H3("crossbar 条件探点与联合路线")
    P("对已获得方向观测的频道，设当前可行多边形为 P，参考测点为 s，读数方向为 a，取纵向单位向量 u 和横向单位向量 v，在 P 沿 u 的中截面两侧布置 crossbar 探点。")
    EQ("u=(cos a, sin a)ᵀ， v=(−sin a, cos a)ᵀ， x±=s+du±hv")
    P("d 取投影区间中点，h 由测向误差宽度和 crossbar_offset=0.03 共同确定；执行前验证两点到当前可行域各顶点的最远距离不超过接收保证。若两侧都返回 no_signal，才用对应截面约束收缩可行域。只要任一点获得正观测，就转入常规定位或清除。该双侧条件把方向歧义转成可核验的几何切割。")
    P("Q4 还使用 40 m 试探清除、每频道至多 12 次无线测量、失败清除圆盘的负区域更新，以及源任务对 21 站骨干的精确插入。局部动作只改变路线次序，CERTIFIED_ABSENT 和全局停止仍由完整证书与基数规则决定。")
    P("21 站证书骨干与局部定位、清除动作的动态插入见图 15。")

    H3("Q4 配对结果")
    TCAP("表 6  Q4 严格配对结果")
    _add_table_before(doc, anchor,
                      ["源数", "案例对数", "V1/s·源⁻¹", "V2/s·源⁻¹", "平均降幅", "V2 更快"],
                      q4_table)
    P(f'Q4 的 60 对 mixed 案例中，V1 和 V2 的平均每源时间分别为 {float(q4_all["mean_v1_s"]):.2f} s 和 {float(q4_all["mean_v2_s"]):.2f} s；平均降低 {float(q4_all["relative_reduction_pct"]):.1f}%，95% bootstrap CI 为 {float(q4_all["relative_reduction_ci95_low_pct"]):.1f}%～{float(q4_all["relative_reduction_ci95_high_pct"]):.1f}%，{q4_all["wins_v2"]}/{q4_all["n_pairs"]} 个案例更快。逐例结果见图 9。')

    H3("Q4 消融、安全门禁与敏感性")
    P("去除 crossbar 后，10 源和 16 源案例的每源时间分别上升 12.9% 和 28.5%，它是 Q4 的主要效率模块。去除精确源插入、负区域更新、初始旋转选择或试探清除时，平均变化不超过 1%，这些模块在本组案例中只有条件性或次要贡献。")
    P("去除 16 源基数停止后，10 源案例仍 20/20 通过，但 16 源案例只有 3/20 通过完整 verifier。其余 17 例虽然清除了全部 16 个真实源，却留下无法证空的频道；这些运行不进入效率比较，不能把不完整终止产生的低耗时写成性能改进。")
    PIC("fig11_q4_ablation.jpg")
    CAP("图 11  Q4 模块消融与 16 源基数停止的安全性负对照")
    P("Q4 的事后邻域分析同样不用于重新选参。无线测量上限 8/12/16 时均为 461.3 s/源；crossbar 偏移 0.02/0.03/0.05 时为 462.9/461.3/459.1 s/源；crossbar 分数 0.15/0.25/0.40 时为 463.6/461.3/467.2 s/源；旋转候选数 6/12/18 时为 464.6/461.3/467.9 s/源；试探半径 20/40/60 m 时为 465.0/461.3/460.1 s/源。全部参数点通过 verifier。")
    PIC("fig12_sensitivity.jpg")
    CAP("图 12  锁定参数邻域的事后敏感性分析")

    H3("时间构成与证据边界")
    P("图 13 将 60 个案例的平均每源虚拟时间分解为移动、测量、换频和清除。主要下降来自移动与测量成本同步减少，与联合开放路线的机制一致。节时并非通过删减验证动作获得：两问共 240/240 次主运行均通过完整门禁。")
    PIC("fig13_time_breakdown.jpg")
    CAP("图 13  V1 与 V2 平均每源虚拟时间的成本分解")

    H3("代表路线与机制解释")
    P("代表案例只用于解释动作怎样被合并，不替代总体统计。Q3 中，搜索站、定位点和清除点共用一条开放路线；Q4 中，局部定位与清除任务插入 21 站证书骨干。图中的真实源均只用于离线核验。")
    PIC("fig14_q3_route.jpg", 5.55)
    CAP("图 14  Q3 代表案例中的联合开放路线")
    PIC("fig15_q4_route.jpg", 5.55)
    CAP("图 15  Q4 代表案例中的 21 站骨干与局部服务插入")


def _replace_validation_and_evaluation(doc: Document) -> None:
    correct = _find_paragraph(doc, "正确性与独立验证")
    robust = _find_paragraph(doc, "稳健性检验")
    evaluation = _find_paragraph(doc, "模型评价与推广")
    _remove_between(correct, robust)
    _add_paragraph_before(
        doc, robust,
        "主实验包含两问各 60 个共同案例、每案例两个策略，共 240 次运行。每次运行均保存案例哈希、代码哈希、参数指纹、动作日志、ground truth、引擎结果、计时指标与 verifier 报告。接受门槛同时要求正常结束、真实源全部清除、世界状态一致、完整 verifier 通过、源代码运行前后哈希一致和虚拟时间账本闭合；240/240 次主运行满足全部条件。V1 与 V2 共享案例文件和计时规则。",
    )
    _remove_between(robust, evaluation)
    _add_paragraph_before(
        doc, evaluation,
        "分层结果在 10、13、16 个源上方向一致，且两问所有配对点都位于等耗时线下方。bootstrap 区间反映当前案例总体中的配对变异，不是确定性最坏情形界。参数敏感性只检查锁定值附近是否出现灾难性变化；即使个别邻近值在 10 个事后案例上略好，也不据此修改冻结参数。消融中不通过 verifier 的运行仅报告安全性通过率。",
    )

    q34_adv = _find_prefix(doc, "问题三、问题四用集合知识状态")
    q34_adv.text = (
        "问题三、问题四将 V1 固定为可执行的状态—证书—停止基线，V2 在相同门禁下优化"
        "动作次序与共享路线。严格配对、源数分层、bootstrap 区间、逐例胜负和动作级证据"
        "共同支持效率结论；Q4 基数停止的失败负对照还能识别由不完整终止造成的伪优势。"
    )
    _format_body(q34_adv)
    q34_lim = _find_prefix(doc, "问题三、问题四以最坏情形保证")
    q34_lim.text = (
        "问题三、问题四的经验结果来自合成随机/混合案例，官方隐藏分布未确认，不能外推为"
        "所有实例上的全局最优。Q3 的若干条件模块边际效应较小；Q4 的 21 站网络是连续覆盖"
        "的充分构造，不声称站点数最少。若实际平台限制目标圆外测量，必须重新设计并认证站网。"
    )
    _format_body(q34_lim)


def _complete_delivery_sections(doc: Document) -> None:
    statement = _find_prefix(doc, "本参赛队在竞赛过程中使用了AI工具")
    statement.text = (
        "本参赛队在竞赛过程中使用了 AI 工具，主要用于语言校核、代码调试、实验脚本检查、"
        "图表生成与文档排版。模型选择、参数锁定、实验运行和结果核验均保留可追溯记录，"
        "详细使用情况与证据文件见支撑材料。"
    )
    _format_body(statement)

    support_rows = [
        ("paper_evidence/cases/", "Q3/Q4 共同案例文件、种子和源数分层"),
        ("paper_evidence/tables/", "主比较、配对效应、消融和敏感性汇总表"),
        ("paper_evidence/figures/", "论文图 7～15 的 RGB 高分辨率 JPG"),
        ("paper_pipeline/ 与 paper_evidence/*.json", "复现实验脚本、来源清单、参数注册表和哈希报告"),
    ]
    appendix_names = [
        "V1 基线叙述与冻结配置",
        "V2 最终代码与锁定参数",
        "主实验严格配对证据",
        "Q3/Q4 模块消融证据",
        "锁定参数邻域敏感性",
        "绘图数据、JPG 与图件清单",
        "verifier、计时账本与复现说明",
    ]
    appendix_details = [
        "列出 V1 文档来源、冻结配置及其状态—证书—终止逻辑。",
        "列出 Q3 v5、Q4 v4 最终求解器的代码哈希、版本和完整参数指纹。",
        "包含 120 个共同案例、240 次主运行的动作日志、ground truth、指标和配对区间。",
        "包含 Q3 5 个模块、Q4 6 个模块及完整策略的分层消融结果；失败局只计安全性。",
        "包含 9 个参数、每参数 3 个邻域值的事后敏感性结果，不用于重新锁参。",
        "包含图 7～15 的源数据、生成脚本、RGB JPG、尺寸与 SHA-256 清单。",
        "包含独立 verifier 报告、世界状态一致性、计时账本闭合和已知测试限制。",
    ]

    support_table = next(
        table for table in doc.tables
        if table.rows and table.rows[0].cells[0].text.strip() == "支撑文件名称"
    )
    for row, (name, description) in zip(support_table.rows[1:], support_rows):
        row.cells[0].text = name
        row.cells[1].text = description

    directory = next(
        table for table in doc.tables
        if len(table.rows) == 8 and table.rows[0].cells[0].text.strip() == "附录"
    )
    for index, name in enumerate(appendix_names, start=1):
        directory.rows[index].cells[1].text = name

    detail_tables = [
        table for table in doc.tables
        if len(table.columns) == 1 and len(table.rows) == 3
        and table.rows[0].cells[0].text.strip().startswith("附录")
    ]
    for table, detail in zip(detail_tables, appendix_details):
        table.rows[2].cells[0].text = detail

    for table in (support_table, directory, *detail_tables):
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.first_line_indent = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    for run in paragraph.runs:
                        _set_run_font(run, size=9, bold=(row_index == 0))

    q3_start = _find_paragraph(doc, "问题三模型建立与求解")
    after_q3 = False
    for paragraph in doc.paragraphs:
        if paragraph._p is q3_start._p:
            after_q3 = True
        if after_q3 and paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True


def update(input_path: Path, output_path: Path, evidence: Path) -> dict:
    doc = Document(str(input_path))
    q12_analysis_before = _range_hash(doc, "问题一：", "问题三：")
    q12_model_before = _range_hash(doc, "问题一模型的建立与求解", "问题三模型建立与求解")

    _replace_short_sections(doc)
    _update_q3_q4_symbols(doc)
    start = _find_paragraph(doc, "问题三模型建立与求解")
    end = _find_paragraph(doc, "模型分析检验")
    _remove_between(start, end, include_start=True)
    _build_model_sections(doc, end, evidence)
    _replace_validation_and_evaluation(doc)
    _complete_delivery_sections(doc)

    q12_analysis_after = _range_hash(doc, "问题一：", "问题三：")
    q12_model_after = _range_hash(doc, "问题一模型的建立与求解", "问题三模型建立与求解")
    if q12_analysis_before != q12_analysis_after:
        raise RuntimeError("Q1/Q2 analysis XML changed unexpectedly")
    if q12_model_before != q12_model_after:
        raise RuntimeError("Q1/Q2 model XML changed unexpectedly")

    doc.core_properties.title = "有界误差下机器狗干扰源定位、选点与搜索清除模型"
    doc.core_properties.comments = "Q3/Q4 replaced with verified V1 baseline + V2 final evidence"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    report = {
        "input": str(input_path),
        "output": str(output_path),
        "input_sha256": hashlib.sha256(input_path.read_bytes()).hexdigest(),
        "output_sha256": hashlib.sha256(output_path.read_bytes()).hexdigest(),
        "q1_q2_analysis_sha256": q12_analysis_after,
        "q1_q2_model_sha256": q12_model_after,
        "q3_q4_symbol_rows_updated": 9,
        "figures_embedded": 9,
    }
    return report


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--evidence", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()
    report = update(args.input.resolve(), args.output.resolve(), args.evidence.resolve())
    args.report.resolve().write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
